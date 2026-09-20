from pathlib import Path


def read_file(file_path: str) -> str:
    """
    Read supported files safely from the project's storage/data folders.

    Supported:
        .txt
        .md
        .csv
        .pdf
        .docx
        .pptx
        .xlsx
        .xls
    """

    try:
        # =========================================================
        # RESOLVE FILE PATH
        # =========================================================
        path = Path(file_path).resolve()

        # =========================================================
        # PROJECT ROOT
        # =========================================================
        project_root = Path(__file__).resolve().parents[3]

        # =========================================================
        # ALLOWED DIRECTORIES
        # =========================================================
        allowed_dirs = [
            (project_root / "storage").resolve(),
            (project_root / "data").resolve(),
        ]

        # =========================================================
        # SECURITY CHECK
        # =========================================================
        if not any(
            path == allowed_dir or allowed_dir in path.parents
            for allowed_dir in allowed_dirs
        ):
            return (
                "Error: Reading outside the allowed project "
                "folders is not allowed."
            )

        # =========================================================
        # FILE EXISTENCE CHECK
        # =========================================================
        if not path.exists():
            return "Error: File does not exist."

        if not path.is_file():
            return "Error: Path is not a file."

        # =========================================================
        # SUPPORTED FILE EXTENSIONS
        # =========================================================
        supported_extensions = [
            ".txt",
            ".md",
            ".csv",
            ".pdf",
            ".docx",
            ".pptx",
            ".xlsx",
            ".xls",
        ]

        extension = path.suffix.lower()

        if extension not in supported_extensions:
            return "Error: File type not supported yet."

        # =========================================================
        # PDF
        # =========================================================
        if extension == ".pdf":

            try:
                from pypdf import PdfReader

                reader = PdfReader(str(path))

                text_parts = []

                for page in reader.pages:
                    text = page.extract_text() or ""

                    if text.strip():
                        text_parts.append(text.strip())

                extracted_text = "\n\n".join(text_parts).strip()

                if not extracted_text:
                    return (
                        "Error: No readable text was found in this PDF. "
                        "The PDF may be scanned or image-based."
                    )

                return extracted_text

            except ImportError:
                return (
                    "Error: pypdf is not installed. "
                    "Run: pip install pypdf"
                )

            except Exception as e:
                return f"PDF reading error: {str(e)}"

        # =========================================================
        # DOCX
        # =========================================================
        if extension == ".docx":

            try:
                from docx import Document

                document = Document(str(path))

                text_parts = []

                # -------------------------------------------------
                # READ PARAGRAPHS
                # -------------------------------------------------
                for paragraph in document.paragraphs:

                    text = paragraph.text.strip()

                    if text:
                        text_parts.append(text)

                # -------------------------------------------------
                # READ TABLES
                # -------------------------------------------------
                for table in document.tables:

                    for row in table.rows:

                        row_data = []

                        for cell in row.cells:

                            cell_text = cell.text.strip()

                            if cell_text:
                                row_data.append(cell_text)

                        if row_data:
                            text_parts.append(
                                " | ".join(row_data)
                            )

                extracted_text = "\n".join(text_parts).strip()

                if not extracted_text:
                    return (
                        "Error: No readable text was found "
                        "in this DOCX file."
                    )

                return extracted_text

            except ImportError:
                return (
                    "Error: python-docx is not installed. "
                    "Run: pip install python-docx"
                )

            except Exception as e:
                return f"DOCX reading error: {str(e)}"

        # =========================================================
        # PPTX
        # =========================================================
        if extension == ".pptx":

            try:
                from pptx import Presentation

                presentation = Presentation(str(path))

                text_parts = []

                # -------------------------------------------------
                # READ EVERY SLIDE
                # -------------------------------------------------
                for slide_number, slide in enumerate(
                    presentation.slides,
                    start=1
                ):

                    slide_text = []

                    for shape in slide.shapes:

                        if hasattr(shape, "text"):

                            text = shape.text.strip()

                            if text:
                                slide_text.append(text)

                    if slide_text:

                        slide_content = (
                            f"Slide {slide_number}\n"
                            + "\n".join(slide_text)
                        )

                        text_parts.append(slide_content)

                extracted_text = "\n\n".join(
                    text_parts
                ).strip()

                if not extracted_text:
                    return (
                        "Error: No readable text was found "
                        "in this PPTX file."
                    )

                return extracted_text

            except ImportError:
                return (
                    "Error: python-pptx is not installed. "
                    "Run: pip install python-pptx"
                )

            except Exception as e:
                return f"PPTX reading error: {str(e)}"

        # =========================================================
        # EXCEL - XLSX / XLS
        # =========================================================
        if extension in [".xlsx", ".xls"]:

            try:
                import pandas as pd

                # -------------------------------------------------
                # READ ALL WORKSHEETS
                # -------------------------------------------------
                excel_file = pd.ExcelFile(str(path))

                sheet_names = excel_file.sheet_names

                if not sheet_names:
                    return "Error: No worksheets found in this Excel file."

                # -------------------------------------------------
                # STORE VALID SHEETS
                # -------------------------------------------------
                valid_sheets = []

                for sheet_name in sheet_names:

                    try:
                        df = pd.read_excel(
                            str(path),
                            sheet_name=sheet_name
                        )

                        # Remove completely empty rows
                        df = df.dropna(
                            axis=0,
                            how="all"
                        )

                        # Remove completely empty columns
                        df = df.dropna(
                            axis=1,
                            how="all"
                        )

                        if df.empty:
                            continue

                        # Count useful/non-empty cells
                        non_empty_cells = (
                            df.notna().sum().sum()
                        )

                        if non_empty_cells == 0:
                            continue

                        valid_sheets.append(
                            {
                                "name": sheet_name,
                                "data": df,
                                "cells": non_empty_cells,
                            }
                        )

                    except Exception:
                        # Ignore a worksheet that cannot be read
                        continue

                if not valid_sheets:
                    return (
                        "Error: No readable data was found "
                        "in this Excel file."
                    )

                # -------------------------------------------------
                # FIND THE BEST DATA SHEET
                #
                # We do NOT simply select the first worksheet.
                #
                # This is important for files exported from
                # Apple Numbers because the first worksheet can
                # be an Export Summary rather than the real data.
                # -------------------------------------------------
                best_sheet = None
                best_score = -1

                for sheet in valid_sheets:

                    sheet_name = sheet["name"]
                    df = sheet["data"]

                    rows, columns = df.shape

                    score = 0

                    # ---------------------------------------------
                    # Prefer larger datasets
                    # ---------------------------------------------
                    score += min(rows, 5000) * 2
                    score += min(columns, 100) * 5

                    # ---------------------------------------------
                    # Penalize obvious metadata/pivot sheets
                    # ---------------------------------------------
                    lower_name = sheet_name.lower()

                    if "export summary" in lower_name:
                        score -= 10000

                    if "pivot" in lower_name:
                        score -= 5000

                    # ---------------------------------------------
                    # Penalize unnamed columns
                    # ---------------------------------------------
                    unnamed_count = sum(
                        1
                        for column in df.columns
                        if str(column).lower().startswith("unnamed")
                    )

                    score -= unnamed_count * 500

                    # ---------------------------------------------
                    # Prefer sheets with more columns
                    # and substantial rows
                    # ---------------------------------------------
                    if rows >= 20:
                        score += 1000

                    if columns >= 5:
                        score += 500

                    # ---------------------------------------------
                    # Select highest score
                    # ---------------------------------------------
                    if score > best_score:

                        best_score = score
                        best_sheet = sheet

                # -------------------------------------------------
                # SAFETY CHECK
                # -------------------------------------------------
                if best_sheet is None:
                    return (
                        "Error: Could not identify a useful "
                        "worksheet in this Excel file."
                    )

                selected_sheet_name = best_sheet["name"]
                df = best_sheet["data"].copy()

                # -------------------------------------------------
                # CLEAN COLUMN NAMES
                # -------------------------------------------------
                cleaned_columns = []

                for column in df.columns:

                    column_name = str(column).strip()

                    if (
                        not column_name
                        or column_name.lower().startswith("unnamed")
                    ):
                        column_name = "Unnamed Column"

                    cleaned_columns.append(column_name)

                df.columns = cleaned_columns

                # -------------------------------------------------
                # REMOVE COMPLETELY EMPTY ROWS AGAIN
                # -------------------------------------------------
                df = df.dropna(
                    axis=0,
                    how="all"
                )

                # -------------------------------------------------
                # LIMIT EXTREMELY LARGE DATASETS
                #
                # This prevents sending millions of rows to
                # the LLM.
                # -------------------------------------------------
                MAX_ROWS = 5000

                truncated = False

                if len(df) > MAX_ROWS:

                    df = df.head(MAX_ROWS)

                    truncated = True

                # -------------------------------------------------
                # CONVERT DATAFRAME TO LLM-FRIENDLY TEXT
                # -------------------------------------------------
                output_parts = []

                output_parts.append(
                    "=== SPREADSHEET ==="
                )

                output_parts.append(
                    f"Worksheet: {selected_sheet_name}"
                )

                output_parts.append(
                    f"Rows: {len(df)}"
                )

                output_parts.append(
                    f"Columns: {len(df.columns)}"
                )

                if truncated:

                    output_parts.append(
                        f"Note: Only the first {MAX_ROWS} rows "
                        "were loaded."
                    )

                output_parts.append(
                    "\nColumn names:"
                )

                output_parts.append(
                    ", ".join(
                        str(column)
                        for column in df.columns
                    )
                )

                output_parts.append(
                    "\n=== DATA ==="
                )

                # -------------------------------------------------
                # CONVERT NaN TO EMPTY VALUES
                # -------------------------------------------------
                display_df = df.copy()

                display_df = display_df.fillna("")

                # -------------------------------------------------
                # Convert rows into readable records
                # -------------------------------------------------
                for index, row in display_df.iterrows():

                    row_values = []

                    for column in display_df.columns:

                        value = row[column]

                        # Convert pandas/numpy values to string
                        if pd.isna(value):
                            value = ""

                        row_values.append(
                            f"{column}: {value}"
                        )

                    output_parts.append(
                        f"Row {index + 1}: "
                        + " | ".join(row_values)
                    )

                # -------------------------------------------------
                # ADD WORKSHEET INFORMATION
                # -------------------------------------------------
                output_parts.append(
                    "\n=== WORKSHEETS FOUND ==="
                )

                for sheet in valid_sheets:

                    output_parts.append(
                        f"- {sheet['name']}: "
                        f"{sheet['data'].shape[0]} rows × "
                        f"{sheet['data'].shape[1]} columns"
                    )

                return "\n".join(output_parts).strip()

            except ImportError:
                return (
                    "Error: pandas is not installed. "
                    "Run: pip install pandas openpyxl"
                )

            except Exception as e:
                return f"Excel reading error: {str(e)}"

        # =========================================================
        # TXT / MD / CSV
        # =========================================================
        if extension in [".txt", ".md", ".csv"]:

            try:

                return path.read_text(
                    encoding="utf-8",
                    errors="ignore"
                ).strip()

            except Exception as e:

                return f"Text file reading error: {str(e)}"

        # =========================================================
        # FALLBACK
        # =========================================================
        return "Error: File type not supported yet."

    # =============================================================
    # GENERAL ERROR HANDLER
    # =============================================================
    except Exception as e:

        return f"File reading error: {str(e)}"